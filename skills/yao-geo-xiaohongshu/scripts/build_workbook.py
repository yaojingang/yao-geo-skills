"""Validate Xiaohongshu autocomplete captures and export a two-level keyword workbook."""
from __future__ import annotations
import argparse
import json
import os
from datetime import datetime
from pathlib import Path
from tempfile import NamedTemporaryFile
from urllib.parse import urlparse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from score_keywords import ALIASES, load_profile, score_keyword

STATUSES = {'ok', 'empty', 'blocked', 'error', 'pending'}

def clean(value, label):
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f'{label} 必须是非空文本')
    return value.strip()

def safe_cell(value):
    if isinstance(value, str) and value.startswith(('=', '+', '-', '@')):
        return "'" + value
    return value

def evidence_path(value, base):
    p = Path(clean(value, '截图路径'))
    if p.is_absolute() or '..' in p.parts or p.suffix.lower() not in {'.png', '.jpg', '.jpeg', '.webp'}:
        raise ValueError(f'无效截图路径: {value}')
    resolved = (base / p).resolve()
    if not resolved.is_relative_to(base.resolve()) or not resolved.is_file():
        raise ValueError(f'截图不存在: {value}')
    return p.as_posix()

def load_capture(path):
    data = json.loads(path.read_text(encoding='utf-8'))
    if data.get('platform') != 'xiaohongshu':
        raise ValueError('platform 必须是 xiaohongshu')
    roots = data.get('roots')
    if not isinstance(roots, list) or not roots:
        raise ValueError('roots 必须是非空数组')
    roots = list(dict.fromkeys(clean(x, '根词') for x in roots))
    clean(data.get('industry'), '行业')
    clean(data.get('conversion_goal'), '转化目标')
    queries = data.get('queries')
    if not isinstance(queries, list):
        raise ValueError('queries 必须是数组')
    by_term = {}
    for i, q in enumerate(queries, 1):
        if not isinstance(q, dict):
            raise ValueError(f'queries[{i}] 必须是对象')
        term = clean(q.get('query'), '检索词')
        if term in by_term:
            raise ValueError(f'检索词重复采集: {term}')
        if q.get('status') not in STATUSES:
            raise ValueError(f'{term}: 无效状态')
        depth = q.get('depth')
        if type(depth) is not int or depth not in (0, 1):
            raise ValueError(f'{term}: depth 必须为 0 或 1')
        if term in roots and depth != 0:
            raise ValueError(f'{term}: 词根的 depth 必须为 0')
        if term not in roots and depth != 1:
            raise ValueError(f'{term}: 非词根的 depth 必须为 1')
        if q['status'] != 'pending':
            captured_at = clean(q.get('captured_at'), f'{term}: 采集时间')
            try:
                parsed_at = datetime.fromisoformat(captured_at.replace('Z', '+00:00'))
            except ValueError as exc:
                raise ValueError(f'{term}: 采集时间不是 ISO 8601') from exc
            if 'T' not in captured_at or parsed_at.utcoffset() is None:
                raise ValueError(f'{term}: 采集时间需包含时区')
        url = q.get('page_url', '')
        if q['status'] in {'ok', 'empty'} and not url:
            raise ValueError(f'{term}: 已访问页面需记录 URL')
        if url and (not isinstance(url, str) or urlparse(url).hostname not in {'xiaohongshu.com', 'www.xiaohongshu.com'} or urlparse(url).scheme != 'https'):
            raise ValueError(f'{term}: 页面 URL 不是小红书')
        shots = q.get('screenshots', [])
        if not isinstance(shots, list):
            raise ValueError(f'{term}: screenshots 必须是数组')
        q['screenshots'] = [evidence_path(s, path.parent) for s in shots]
        suggestions = q.get('suggestions', [])
        if not isinstance(suggestions, list):
            raise ValueError(f'{term}: suggestions 必须是数组')
        seen_rank, seen_term = set(), set()
        for s in suggestions:
            if not isinstance(s, dict):
                raise ValueError(f'{term}: 无效建议词')
            word = clean(s.get('term'), '相关词')
            rank = s.get('rank')
            if type(rank) is not int or rank < 1 or rank in seen_rank or word in seen_term:
                raise ValueError(f'{term}: 下拉位置或建议词重复')
            seen_rank.add(rank); seen_term.add(word)
        if q['status'] == 'ok' and (not suggestions or not shots):
            raise ValueError(f'{term}: ok 状态需要建议词和截图')
        if q['status'] == 'empty' and not shots:
            raise ValueError(f'{term}: empty 状态需要空列表截图')
        if q['status'] != 'ok' and suggestions:
            raise ValueError(f'{term}: 非 ok 状态不能含建议词')
        if q['status'] in {'blocked', 'error'}:
            clean(q.get('note'), f'{term}: 受阻或报错原因')
        by_term[term] = q
    return data, roots, by_term

def make_rows(roots, by_term, profile):
    rows = []
    for root in roots:
        q = by_term.get(root)
        if not q or q['status'] != 'ok':
            continue
        for first in sorted(q['suggestions'], key=lambda s:s['rank']):
            term = first['term'].strip()
            score, reason = score_keyword(term, profile)
            rows.append([root, 1, root, term, first['rank'], score, '', '', q.get('captured_at',''), q.get('page_url',''), q['screenshots'][0], reason])
            if term == root:
                continue
            child = by_term.get(term)
            if child and child['status'] == 'ok':
                for second in sorted(child['suggestions'], key=lambda s:s['rank']):
                    word = second['term'].strip()
                    score, reason = score_keyword(word, profile)
                    rows.append([root, 2, term, word, second['rank'], score, term, first['rank'], child.get('captured_at',''), child.get('page_url',''), child['screenshots'][0], reason])
    return rows

def pending_queries(roots, by_term):
    required = list(roots)
    for root in roots:
        q = by_term.get(root)
        if q and q['status'] == 'ok':
            required.extend(s['term'].strip() for s in sorted(q['suggestions'], key=lambda x:x['rank']))
    incomplete = []
    for term in dict.fromkeys(required):
        q = by_term.get(term)
        if not q or q['status'] in {'pending', 'blocked', 'error'}:
            status = q['status'] if q else 'pending'
            incomplete.append((term, status, q.get('note', '') if q else '尚未检索'))
    return incomplete

def add_sheet(wb, title, headers, rows):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append([safe_cell(v) for v in row])
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 27
    for cell in ws[1]:
        cell.fill = PatternFill('solid', fgColor='123B5D')
        cell.font = Font(color='FFFFFF', bold=True)
        cell.alignment = Alignment(vertical='center')
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        length = max(len(str(c.value or '')) for c in col)
        ws.column_dimensions[letter].width = min(max(length * 1.5, 12), 48)
    return ws

def build(input_path, output_path, industry=None, profile_path=None):
    if input_path.resolve().parent != output_path.resolve().parent:
        raise ValueError('Excel 必须与采集 JSON 放在同一目录，确保截图相对链接可用')
    data, roots, by_term = load_capture(input_path)
    if industry and ALIASES.get(industry, industry) != ALIASES.get(data['industry'], data['industry']):
        raise ValueError('命令行行业与采集 JSON 的行业不一致')
    selected = industry or data.get('industry')
    profile = load_profile(selected, profile_path)
    if data['conversion_goal'].strip() != profile['conversion_goal'].strip():
        raise ValueError('采集 JSON 的转化目标与评分画像不一致')
    rows = make_rows(roots, by_term, profile)
    wb = Workbook(); wb.remove(wb.active)
    ws = add_sheet(wb, '关键词表', ['根词','层级','检索词','相关搜索词','下拉位置','商业价值评分（10分）','一级相关词','一级位置','采集时间','页面URL','截图','评分理由'], rows)
    for i, row in enumerate(rows, 2):
        ws.cell(i, 6).number_format = '0'
        ws.cell(i, 11).hyperlink = row[10]
        ws.cell(i, 11).style = 'Hyperlink'
    records = []
    for term, q in by_term.items():
        records.append([term, q.get('depth',''), q['status'], len(q['suggestions']), q.get('captured_at',''), q.get('page_url',''), ', '.join(q['screenshots']), q.get('note','')])
    add_sheet(wb, '采集记录', ['检索词','查询层级','状态','建议词数','采集时间','页面URL','截图','备注'], records)
    evidence = [[term, q.get('depth', ''), index, path]
                for term, q in by_term.items()
                for index, path in enumerate(q['screenshots'], 1)]
    evidence_ws = add_sheet(wb, '截图索引', ['检索词','查询层级','截图序号','截图路径'], evidence)
    for row_number, item in enumerate(evidence, 2):
        evidence_ws.cell(row_number, 4).hyperlink = item[3]
        evidence_ws.cell(row_number, 4).style = 'Hyperlink'
    unique = {}
    for row in rows:
        term = row[3]
        if term not in unique:
            unique[term] = [term,row[5],row[11],0,set()]
        unique[term][3] += 1; unique[term][4].add(row[0])
    add_sheet(wb, '去重词库', ['相关搜索词','商业价值评分（10分）','评分理由','出现路径数','根词'], [[v[0],v[1],v[2],v[3],'、'.join(sorted(v[4]))] for v in unique.values()])
    pending = pending_queries(roots, by_term)
    add_sheet(wb, '待采集', ['检索词','状态','说明'], pending)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temp_path = None
    try:
        with NamedTemporaryFile(dir=output_path.parent, prefix='.xiaohongshu-', suffix='.xlsx', delete=False) as handle:
            temp_path = Path(handle.name)
        wb.save(temp_path)
        os.replace(temp_path, output_path)
    finally:
        if temp_path is not None:
            temp_path.unlink(missing_ok=True)
    return len(rows), len(unique), len(pending)

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--input', required=True, type=Path)
    ap.add_argument('--output', required=True, type=Path)
    ap.add_argument('--industry')
    ap.add_argument('--profile', type=Path)
    args = ap.parse_args()
    print('路径行数、去重词数、待采集数:', *build(args.input,args.output,args.industry,args.profile))
if __name__ == '__main__': main()
