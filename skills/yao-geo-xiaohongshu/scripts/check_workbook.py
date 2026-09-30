#!/usr/bin/env python3
"""Verify two-level paths, shared-query reuse, pending work, scores, and evidence checks."""
import json
from pathlib import Path
from tempfile import TemporaryDirectory
from openpyxl import load_workbook
from build_workbook import build, load_capture

def main():
    with TemporaryDirectory() as folder:
        base = Path(folder); (base/'evidence').mkdir()
        for name in ('root.png','other.png','child.png'):
            (base/'evidence'/name).write_bytes(b'\x89PNG\r\n\x1a\n')
        data = {
            'platform':'xiaohongshu','industry':'study_abroad',
            'conversion_goal':'咨询或委托留学服务机构进行申请规划与办理',
            'roots':['留学公司','留学机构'],
            'queries':[
                {'query':'留学公司','depth':0,'status':'ok','captured_at':'2026-09-29T10:00:00+08:00','page_url':'https://www.xiaohongshu.com/explore','screenshots':['evidence/root.png'],'suggestions':[{'term':'留学咨询','rank':2},{'term':'留学费用','rank':1}]},
                {'query':'留学机构','depth':0,'status':'ok','captured_at':'2026-09-29T10:00:00+08:00','page_url':'https://www.xiaohongshu.com/explore','screenshots':['evidence/other.png'],'suggestions':[{'term':'留学咨询','rank':1}]},
                {'query':'留学咨询','depth':1,'status':'ok','captured_at':'2026-09-29T10:00:00+08:00','page_url':'https://www.xiaohongshu.com/explore','screenshots':['evidence/child.png'],'suggestions':[{'term':'留学咨询机构','rank':1},{'term':'留学公司','rank':2}]}
            ]}
        path = base/'capture.json'; path.write_text(json.dumps(data,ensure_ascii=False),encoding='utf-8')
        out=base/'keywords.xlsx'
        rows,unique,pending=build(path,out)
        assert (rows,unique,pending)==(7,4,1),(rows,unique,pending)
        wb=load_workbook(out)
        assert wb.sheetnames==['关键词表','采集记录','截图索引','去重词库','待采集']
        assert wb['截图索引'].max_row == 4
        assert wb['截图索引']['D2'].hyperlink.target == 'evidence/root.png'
        values=list(wb['关键词表'].values)
        assert values[1][3:6]==('留学费用',1,values[1][5])
        assert values[2][3:5]==('留学咨询',2)
        assert values[3][1:8]==(2,'留学咨询','留学咨询机构',1,values[3][5],'留学咨询',2)
        assert values[6][0]=='留学机构' and values[6][1]==2
        assert wb['待采集']['A2'].value=='留学费用'
        assert all(type(row[5]) is int and 1<=row[5]<=10 for row in values[1:])
        assert len({row[5] for row in values[1:] if row[3]=='留学咨询机构'})==1
        assert wb['待采集']['B2'].value=='pending'
        data['queries'][0]['suggestions'].append({'term':'留学公司','rank':3})
        path.write_text(json.dumps(data,ensure_ascii=False),encoding='utf-8')
        rows,_,pending=build(path,out)
        assert (rows,pending)==(8,1), (rows,pending)
        data['queries'].append({'query':'留学费用','depth':1,'status':'blocked','captured_at':'2026-09-29T10:00:00+08:00','note':'需要登录','screenshots':[],'suggestions':[]})
        path.write_text(json.dumps(data,ensure_ascii=False),encoding='utf-8')
        rows,_,pending=build(path,out)
        assert (rows,pending)==(8,1)
        assert load_workbook(out)['待采集']['B2'].value=='blocked'
        try: build(path,base/'other'/'keywords.xlsx')
        except ValueError: pass
        else: raise AssertionError('应拒绝截图链接失效的输出目录')
        data['conversion_goal']='销售留学机构加盟服务'
        path.write_text(json.dumps(data,ensure_ascii=False),encoding='utf-8')
        try: build(path,out)
        except ValueError: pass
        else: raise AssertionError('应拒绝转化目标与评分画像不一致')
        data['conversion_goal']='咨询或委托留学服务机构进行申请规划与办理'
        data['queries'][0]['screenshots']=['../outside.png']
        path.write_text(json.dumps(data,ensure_ascii=False),encoding='utf-8')
        try: load_capture(path)
        except ValueError: pass
        else: raise AssertionError('应拒绝越界截图路径')
    print('通过两层路径、共享查询、待采集、评分及截图路径检查')
if __name__=='__main__': main()
