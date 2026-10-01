"""Validate Douyin autocomplete captures and export a two-level keyword workbook."""
from __future__ import annotations
import argparse
import json
import os
import re
from datetime import datetime
from pathlib import Path
from tempfile import NamedTemporaryFile
from urllib.parse import quote, urlparse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from PIL import Image
from score_keywords import ALIASES, load_profile, score_keyword

STATUSES = {'ok', 'empty', 'blocked', 'error', 'pending'}
INVALID_XML_TEXT = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]')

def validate_text(value, label):
    if len(value) > 32767:
        raise ValueError(f'{label} 超过Excel单元格长度，不能无损导出；请拆分批次或汇总字段')
    if INVALID_XML_TEXT.search(value):
        raise ValueError(f'{label} 含Excel XML不支持的字符，不能无损导出')
    return value

def clean(value, label):
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f'{label} 必须是非空文本')
    return value.strip()

def keyword(value, label):
    clean(value, label)
    return validate_text(value, label)

def screenshot_reference(value):
    clean(value, '截图路径')
    if re.search(r'[\x00-\x1f\x7f\\<>:"|?*]', value):
        raise ValueError(f'截图路径必须是包内可跨系统使用的相对路径: {value!r}')
    parts = value.split('/')
    for part in parts:
        if part in {'', '.'}:
            continue
        if part == '..' or part.endswith((' ', '.')):
            raise ValueError(f'无效截图路径分段: {value!r}')
        if re.fullmatch(r'CON|PRN|AUX|NUL|CONIN\$|CONOUT\$|COM[1-9¹²³]|LPT[1-9¹²³]', part.split('.')[0].rstrip(' '), re.IGNORECASE):
            raise ValueError(f'截图路径含系统保留文件名: {value!r}')
    p = Path(value)
    if p.is_absolute() or '..' in p.parts or p.suffix.lower() not in {'.png', '.jpg', '.jpeg', '.webp'}:
        raise ValueError(f'无效截图路径: {value!r}')
    return validate_text(p.as_posix(), '截图路径')

def evidence_path(value, base):
    p = Path(screenshot_reference(value))
    resolved = (base / p).resolve()
    if not resolved.is_relative_to(base.resolve()) or not resolved.is_file():
        raise ValueError(f'截图不存在: {value}')
    with resolved.open('rb') as handle:
        header = handle.read(16)
    actual = ('png' if header.startswith(b'\x89PNG\r\n\x1a\n') else
              'jpeg' if header.startswith(b'\xff\xd8\xff') else
              'webp' if header.startswith(b'RIFF') and header[8:12] == b'WEBP' else None)
    expected = {'.png': 'png', '.jpg': 'jpeg', '.jpeg': 'jpeg', '.webp': 'webp'}[p.suffix.lower()]
    if actual != expected:
        raise ValueError(f'截图实际格式与扩展名不符: {value}')
    try:
        with Image.open(resolved) as picture:
            if picture.width * picture.height > 50_000_000:
                raise ValueError('截图尺寸超过5000万像素')
            picture.verify()
        with Image.open(resolved) as picture:
            picture.load()
    except Exception as exc:
        raise ValueError(f'截图损坏或无法完整解码: {value}: {exc}') from exc
    return p.as_posix()

def load_capture(path):
    data = json.loads(path.read_text(encoding='utf-8'))
    if not isinstance(data, dict):
        raise ValueError('采集 JSON 必须是对象')
    if data.get('platform') != 'douyin':
        raise ValueError('platform 必须是 douyin')
    if data.get('capture_mode') != 'input_only':
        raise ValueError('capture_mode 必须是 input_only')
    roots = data.get('roots')
    if not isinstance(roots, list) or not roots:
        raise ValueError('roots 必须是非空数组')
    roots = list(dict.fromkeys(keyword(x, '根词') for x in roots))
    clean(data.get('industry'), '行业')
    clean(data.get('conversion_goal'), '转化目标')
    queries = data.get('queries')
    if not isinstance(queries, list):
        raise ValueError('queries 必须是数组')
    by_term = {}
    shot_owners = {}
    for i, q in enumerate(queries, 1):
        if not isinstance(q, dict):
            raise ValueError(f'queries[{i}] 必须是对象')
        term = keyword(q.get('query'), '检索词')
        if term in by_term:
            raise ValueError(f'检索词重复采集: {term}')
        if not isinstance(q.get('status'), str) or q['status'] not in STATUSES:
            raise ValueError(f'{term}: 无效状态')
        depth = q.get('depth')
        if type(depth) is not int or depth not in (0, 1):
            raise ValueError(f'{term}: depth 必须为 0 或 1')
        if term in roots and depth != 0:
            raise ValueError(f'{term}: 词根的 depth 必须为 0')
        if term not in roots and depth != 1:
            raise ValueError(f'{term}: 非词根的 depth 必须为 1')
        if q['status'] != 'pending':
            captured_at = clean(q.get('captured_at'), f'{term}: 采集时间')
            try:
                parsed_at = datetime.fromisoformat(captured_at.replace('Z', '+00:00'))
            except ValueError as exc:
                raise ValueError(f'{term}: 采集时间不是 ISO 8601') from exc
            if 'T' not in captured_at or parsed_at.utcoffset() is None:
                raise ValueError(f'{term}: 采集时间需包含时区')
        url = q.get('page_url', '')
        before = q.get('page_url_before', '')
        if q['status'] in {'ok', 'empty'} and not url:
            raise ValueError(f'{term}: 已访问页面需记录 URL')
        for item in (url, before):
            if item and (not isinstance(item, str) or urlparse(item).hostname not in {'douyin.com', 'www.douyin.com'} or urlparse(item).scheme != 'https'):
                raise ValueError(f'{term}: 页面 URL 必须是抖音 HTTPS 页面')
        if q['status'] in {'ok', 'empty'}:
            if not before or before != url:
                raise ValueError(f'{term}: 仅输入模式要求输入前后实际 URL 相同')
            if q.get('observed_input') != term:
                raise ValueError(f'{term}: 输入核对值与检索词不一致')
        if q['status'] == 'empty':
            clean(q.get('empty_confirmation'), f'{term}: 空列表确认')
        if 'list_truncated' in q and type(q['list_truncated']) is not bool:
            raise ValueError(f'{term}: list_truncated 必须为布尔值')
        if q.get('list_truncated'):
            clean(q.get('note'), f'{term}: 截断范围说明')
        shots = q.get('screenshots', [])
        if not isinstance(shots, list):
            raise ValueError(f'{term}: screenshots 必须是数组')
        q['screenshots'] = [evidence_path(s, path.parent) for s in shots]
        for shot in q['screenshots']:
            resolved_shot = (path.parent / shot).resolve()
            file_info = resolved_shot.stat()
            file_identity = (file_info.st_dev, file_info.st_ino)
            if file_identity in shot_owners and shot_owners[file_identity] != term:
                raise ValueError(f'{term}: 不同检索不能共用同一截图文件')
            shot_owners[file_identity] = term
        suggestions = q.get('suggestions', [])
        if not isinstance(suggestions, list):
            raise ValueError(f'{term}: suggestions 必须是数组')
        q['suggestions'] = suggestions
        seen_rank = set()
        for s in suggestions:
            if not isinstance(s, dict):
                raise ValueError(f'{term}: 无效建议词')
            keyword(s.get('term'), '相关词')
            rank = s.get('rank')
            if type(rank) is not int or rank < 1 or rank in seen_rank:
                raise ValueError(f'{term}: 下拉位置无效或重复')
            seen_rank.add(rank)
            shot = s.get('screenshot')
            if shot is None and len(q['screenshots']) == 1:
                shot = q['screenshots'][0]
            elif shot is not None:
                shot = screenshot_reference(shot)
            if shot not in q['screenshots']:
                raise ValueError(f'{term}: 每条建议需对应本查询的截图，多图时不得省略')
            s['screenshot'] = shot
        if q['status'] == 'ok' and (not suggestions or not shots):
            raise ValueError(f'{term}: ok 状态需要建议词和截图')
        if q['status'] == 'empty' and not shots:
            raise ValueError(f'{term}: empty 状态需要空列表截图')
        if q['status'] != 'ok' and suggestions:
            raise ValueError(f'{term}: 非 ok 状态不能含建议词')
        if q['status'] in {'blocked', 'error'}:
            clean(q.get('note'), f'{term}: 受阻或报错原因')
        by_term[term] = q
    first_level = {
        s['term']
        for root in roots
        if root in by_term and by_term[root]['status'] == 'ok'
        for s in by_term[root]['suggestions']
    }
    for term in by_term:
        if term not in roots and term not in first_level:
            raise ValueError(f'{term}: 非根查询必须来自本批次已采集的一级建议')
    return data, roots, by_term

def make_rows(roots, by_term, profile):
    rows = []
    for root in roots:
        q = by_term.get(root)
        if not q or q['status'] != 'ok':
            continue
        for first in sorted(q['suggestions'], key=lambda s:s['rank']):
            term = first['term']
            score, reason = score_keyword(term, profile)
            rows.append([root, 1, root, term, first['rank'], score, '', '', q.get('captured_at',''), q.get('page_url',''), first['screenshot'], reason])
            if term == root:
                continue
            child = by_term.get(term)
            if child and child['status'] == 'ok':
                for second in sorted(child['suggestions'], key=lambda s:s['rank']):
                    word = second['term']
                    score, reason = score_keyword(word, profile)
                    rows.append([root, 2, term, word, second['rank'], score, term, first['rank'], child.get('captured_at',''), child.get('page_url',''), second['screenshot'], reason])
    return rows

def pending_queries(roots, by_term):
    required = list(roots)
    for root in roots:
        q = by_term.get(root)
        if q and q['status'] == 'ok':
            required.extend(s['term'] for s in sorted(q['suggestions'], key=lambda x:x['rank']))
    incomplete = []
    for term in dict.fromkeys(required):
        q = by_term.get(term)
        if not q or q['status'] in {'pending', 'blocked', 'error'}:
            status = q['status'] if q else 'pending'
            incomplete.append((term, status, q.get('note', '') if q else '尚未检索'))
    return incomplete

def add_sheet(wb, title, headers, rows):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for row_number, row in enumerate(rows, 2):
        for column, value in enumerate(row, 1):
            if isinstance(value, str):
                validate_text(value, f'{title}!{get_column_letter(column)}{row_number}')
        ws.append(row)
        for column, value in enumerate(row, 1):
            if isinstance(value, str):
                cell = ws.cell(row_number, column)
                cell.data_type = 's'
                if value.startswith(('=', '+', '-', '@')):
                    cell.quotePrefix = True
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 27
    for cell in ws[1]:
        cell.fill = PatternFill('solid', fgColor='123B5D')
        cell.font = Font(color='FFFFFF', bold=True)
        cell.alignment = Alignment(vertical='center')
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        length = max(len(str(c.value or '')) for c in col)
        ws.column_dimensions[letter].width = min(max(length * 1.5, 12), 48)
    return ws

def build(input_path, output_path, industry=None, profile_path=None):
    if output_path.suffix.lower() != '.xlsx':
        raise ValueError('Excel 输出文件必须使用 .xlsx 扩展名')
    if input_path.resolve() == output_path.resolve():
        raise ValueError('输出文件不能覆盖采集 JSON')
    if input_path.absolute().parent.resolve() != output_path.absolute().parent.resolve():
        raise ValueError('Excel 必须与采集 JSON 放在同一目录，确保截图相对链接可用')
    data, roots, by_term = load_capture(input_path)
    if industry and ALIASES.get(industry, industry) != ALIASES.get(data['industry'], data['industry']):
        raise ValueError('命令行行业与采集 JSON 的行业不一致')
    selected = industry or data.get('industry')
    profile = load_profile(selected, profile_path)
    if data['conversion_goal'].strip() != profile['conversion_goal'].strip():
        raise ValueError('采集 JSON 的转化目标与评分画像不一致')
    rows = make_rows(roots, by_term, profile)
    wb = Workbook(); wb.remove(wb.active)
    ws = add_sheet(wb, '关键词表', ['根词','层级','检索词','相关搜索词','下拉位置','商业价值评分（10分）','一级相关词','一级位置','采集时间','页面URL','截图','评分理由'], rows)
    for i, row in enumerate(rows, 2):
        ws.cell(i, 6).number_format = '0'
        ws.cell(i, 11).hyperlink = quote(row[10], safe='/')
        ws.cell(i, 11).style = 'Hyperlink'
    records = []
    for term, q in by_term.items():
        records.append([term, q.get('depth',''), q['status'], len(q['suggestions']), q.get('captured_at',''), q.get('page_url',''), ', '.join(q['screenshots']), q.get('note',''), q.get('observed_input',''), q.get('page_url_before',''), q.get('list_truncated',False), q.get('empty_confirmation','')])
    add_sheet(wb, '采集记录', ['检索词','查询层级','状态','建议词数','采集时间','页面URL','截图','备注','输入核对值','输入前URL','列表截断','空列表确认'], records)
    evidence = [[term, q.get('depth', ''), index, path]
                for term, q in by_term.items()
                for index, path in enumerate(q['screenshots'], 1)]
    evidence_ws = add_sheet(wb, '截图索引', ['检索词','查询层级','截图序号','截图路径'], evidence)
    for row_number, item in enumerate(evidence, 2):
        evidence_ws.cell(row_number, 4).hyperlink = quote(item[3], safe='/')
        evidence_ws.cell(row_number, 4).style = 'Hyperlink'
    unique = {}
    for row in rows:
        term = row[3]
        if term not in unique:
            unique[term] = [term,row[5],row[11],0,set()]
        unique[term][3] += 1; unique[term][4].add(row[0])
    add_sheet(wb, '去重词库', ['相关搜索词','商业价值评分（10分）','评分理由','出现路径数','根词'], [[v[0],v[1],v[2],v[3],'、'.join(sorted(v[4]))] for v in unique.values()])
    pending = pending_queries(roots, by_term)
    add_sheet(wb, '待采集', ['检索词','状态','说明'], pending)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temp_path = None
    try:
        with NamedTemporaryFile(dir=output_path.parent, prefix='.douyin-', suffix='.xlsx', delete=False) as handle:
            temp_path = Path(handle.name)
        wb.save(temp_path)
        os.replace(temp_path, output_path)
    finally:
        if temp_path is not None:
            temp_path.unlink(missing_ok=True)
    return len(rows), len(unique), len(pending)

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--input', required=True, type=Path)
    ap.add_argument('--output', required=True, type=Path)
    ap.add_argument('--industry')
    ap.add_argument('--profile', type=Path)
    args = ap.parse_args()
    try:
        result = build(args.input, args.output, args.industry, args.profile)
    except (ValueError, OSError) as exc:
        ap.error(str(exc))
    print('路径行数、去重词数、待采集数:', *result)
if __name__ == '__main__': main()
