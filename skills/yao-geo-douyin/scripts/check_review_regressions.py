#!/usr/bin/env python3
"""Check review regressions with synthetic captures, without opening a browser."""
from __future__ import annotations

import copy
import json
import os
import unittest
import zipfile
from pathlib import Path
from tempfile import TemporaryDirectory
from urllib.parse import quote, unquote, urlsplit
from xml.etree import ElementTree

from PIL import Image
from openpyxl import load_workbook

from build_workbook import build
from score_keywords import PROFILE_FILE, load_profile, score_keyword, validate_profile

URL = 'https://www.douyin.com/jingxuan'


def capture_query(term, screenshot, suggestions=None, depth=0):
    return {
        'query': term, 'depth': depth, 'status': 'ok',
        'observed_input': term, 'captured_at': '2026-09-30T21:00:00+08:00',
        'page_url_before': URL, 'page_url': URL,
        'screenshots': [screenshot],
        'suggestions': suggestions or [{'term': term, 'rank': 1}],
        'note': '合成回归样例，不是抖音实采',
    }


class ReviewRegressions(unittest.TestCase):
    def setUp(self):
        self.folder = TemporaryDirectory()
        self.addCleanup(self.folder.cleanup)
        self.base = Path(self.folder.name)
        self.source = self.base / 'capture.json'
        self.output = self.base / 'keywords.xlsx'
        self.picture('evidence/a.png')
        self.data = {
            'platform': 'douyin', 'capture_mode': 'input_only',
            'industry': 'geo_services',
            'conversion_goal': load_profile('geo_services')['conversion_goal'],
            'roots': ['GEO公司'],
            'queries': [capture_query('GEO公司', 'evidence/a.png')],
        }

    def picture(self, relative):
        path = self.base / relative
        path.parent.mkdir(parents=True, exist_ok=True)
        Image.new('RGB', (32, 32), 'white').save(path, format='PNG')

    def save(self, data=None):
        self.source.write_text(json.dumps(data or self.data, ensure_ascii=True), encoding='utf-8')

    @unittest.skipUnless(os.name == 'posix', 'POSIX文件名产生的跨系统边界')
    def test_cross_platform_screenshot_paths(self):
        for value in [r'\\attacker.invalid\share\evidence.png', r'C:\evidence.png',
                      'evidence/https:other.png', 'evidence/.. /outside.png',
                      'evidence/COM¹.png', 'evidence/LPT².png', 'evidence/NUL .png',
                      'evidence/CONIN$.png', 'evidence/CONOUT$.png']:
            with self.subTest(path=value):
                self.picture(value)
                data = copy.deepcopy(self.data)
                data['queries'][0]['screenshots'] = [value]
                self.save(data)
                with self.assertRaises(ValueError):
                    build(self.source, self.output)

    def test_special_filename_link(self):
        relative = 'evidence/采集 #1% 完整.png'
        self.picture(relative)
        self.data['queries'][0]['screenshots'] = [relative]
        self.save()
        build(self.source, self.output)
        wb = load_workbook(self.output)
        for sheet, address in [('关键词表', 'K2'), ('截图索引', 'D2')]:
            cell = wb[sheet][address]
            self.assertEqual(cell.value, relative)
            self.assertEqual(cell.hyperlink.target, quote(relative, safe='/'))
            parsed = urlsplit(cell.hyperlink.target)
            self.assertFalse(parsed.scheme or parsed.netloc or parsed.fragment)
            self.assertTrue((self.base / unquote(parsed.path)).is_file())

    def test_equivalent_screenshot_references(self):
        self.picture('evidence/b.png')
        q = self.data['queries'][0]
        q['screenshots'] = ['./evidence/a.png', './evidence/b.png']
        q['suggestions'][0]['screenshot'] = './evidence/b.png'
        self.save()
        self.assertEqual(build(self.source, self.output), (1, 1, 0))
        self.assertEqual(load_workbook(self.output)['关键词表']['K2'].hyperlink.target,
                         'evidence/b.png')

    def test_different_queries_cannot_reuse_file_aliases(self):
        aliases = []
        try:
            os.link(self.base / 'evidence/a.png', self.base / 'evidence/linked.png')
            aliases.append('evidence/linked.png')
        except OSError:
            pass
        case_alias = self.base / 'evidence/A.PNG'
        if case_alias.exists() and case_alias.samefile(self.base / 'evidence/a.png'):
            aliases.append('evidence/A.PNG')
        if not aliases:
            self.skipTest('当前文件系统不支持文件别名测试')
        for alias in aliases:
            with self.subTest(alias=alias):
                data = copy.deepcopy(self.data)
                data['roots'].append('GEO服务商')
                data['queries'].append(capture_query('GEO服务商', alias))
                self.save(data)
                with self.assertRaises(ValueError):
                    build(self.source, self.output)

    def test_keyword_original_text_and_formula_guard(self):
        terms = ['=1+1', '+律师', '-律师', '@律师', '=HYPERLINK("https://invalid","x")']
        self.data['roots'] = terms
        self.data['queries'] = []
        for i, term in enumerate(terms):
            shot = f'evidence/{i}.png'
            self.picture(shot)
            self.data['queries'].append(capture_query(term, shot))
        self.save()
        self.assertEqual(build(self.source, self.output), (len(terms), len(terms), 0))
        wb = load_workbook(self.output)
        for row, term in enumerate(terms, 2):
            for column in ('A', 'C', 'D'):
                cell = wb['关键词表'][f'{column}{row}']
                self.assertEqual(cell.value, term)
                self.assertEqual(cell.data_type, 's')
        self.assertEqual(set(row[0] for row in list(wb['去重词库'].values)[1:]), set(terms))
        with zipfile.ZipFile(self.output) as archive:
            for name in archive.namelist():
                if name.startswith('xl/worksheets/') and name.endswith('.xml'):
                    tree = ElementTree.fromstring(archive.read(name))
                    self.assertFalse(tree.findall('.//{*}f'))

    def test_industry_names_round_trip(self):
        profiles = json.loads(PROFILE_FILE.read_text(encoding='utf-8'))['profiles']
        for key, original in profiles.items():
            with self.subTest(industry=key):
                self.assertEqual(load_profile(original['industry'])['industry'], original['industry'])
                self.data['industry'] = original['industry']
                self.data['conversion_goal'] = original['conversion_goal']
                self.data['queries'] = [{'query': 'GEO公司', 'depth': 0, 'status': 'pending'}]
                self.save()
                self.assertEqual(build(self.source, self.output), (0, 0, 1))
                self.assertEqual(build(self.source, self.output, industry=key), (0, 0, 1))

    def test_aggregate_text_length_is_not_truncated(self):
        roots = ['留学' + '甲' * 20000, '留学' + '乙' * 20000]
        self.data['roots'] = roots
        self.data['queries'] = []
        for i, root in enumerate(roots):
            shot = f'evidence/large-{i}.png'
            self.picture(shot)
            self.data['queries'].append(capture_query(root, shot, [{'term': '留学咨询', 'rank': 1}]))
        self.save()
        self.output.write_bytes(b'previous output')
        with self.assertRaises(ValueError):
            build(self.source, self.output)
        self.assertEqual(self.output.read_bytes(), b'previous output')

    def test_invalid_xml_text_keeps_previous_output(self):
        for bad in ['留学\ufffe公司', '留学\ud800公司', '留学\x00公司']:
            with self.subTest(value=repr(bad)):
                data = copy.deepcopy(self.data)
                data['queries'][0]['suggestions'][0]['term'] = bad
                self.save(data)
                self.output.write_bytes(b'previous output')
                with self.assertRaises(ValueError):
                    build(self.source, self.output)
                self.assertEqual(self.output.read_bytes(), b'previous output')
        self.data['queries'][0]['note'] = '备注\uffff'
        self.save()
        with self.assertRaises(ValueError):
            build(self.source, self.output)

    def test_json_file_symlink_uses_access_directory(self):
        self.save()
        alias = self.base / 'alias'
        alias.mkdir()
        linked_input = alias / 'capture.json'
        try:
            linked_input.symlink_to(self.source)
        except OSError:
            self.skipTest('当前环境不支持文件符号链接')
        (alias / 'evidence').mkdir()
        Image.new('RGB', (32, 32), 'white').save(alias / 'evidence/a.png')
        with self.assertRaises(ValueError):
            build(linked_input, self.output)
        aliased_output = alias / 'keywords.xlsx'
        self.assertEqual(build(linked_input, aliased_output), (1, 1, 0))
        target = load_workbook(aliased_output)['关键词表']['K2'].hyperlink.target
        self.assertTrue((aliased_output.parent / unquote(target)).is_file())

    def test_information_and_after_sales_scoring(self):
        samples = [
            ('legal', '律师是什么意思', 2),
            ('legal', '律师法全文下载', 3),
            ('legal', '律师起诉状范文下载', 3),
            ('legal', '找律师代写起诉状多少钱', 10),
            ('ecommerce', '扫地机器人官方旗舰店退货', 5),
            ('ecommerce', '扫地机器人官方旗舰店售后电话', 5),
            ('ecommerce', '扫地机器人官方旗舰店维修', 5),
            ('ecommerce', '扫地机器人官方旗舰店购买', 10),
        ]
        for industry, term, expected in samples:
            with self.subTest(term=term):
                self.assertEqual(score_keyword(term, load_profile(industry))[0], expected)

    def test_service_context_and_operator_exclusions(self):
        samples = [
            ('education', '公务员面试培训机构报名', 10),
            ('education', '培训机构面试经验', 2),
            ('education', '面试培训机构招聘老师', 2),
            ('study_abroad', '留学机构面试辅导咨询', 9),
            ('study_abroad', '留学机构面试经验', 2),
            ('study_abroad', '留学机构面试辅导师招聘', 2),
            ('study_abroad', '留学申请辅导机构面试收费', 8),
            ('geo_services', 'GEO优化获客服务商报价', 9),
            ('geo_services', 'GEO服务商怎么获客', 2),
            ('geo_services', '找有GEO优化需求的公司', 2),
            ('legal', '找律师审核合同模板多少钱', 10),
            ('legal', '审核合同模板找律师多少钱', 10),
            ('legal', '找律师起诉状范文下载', 3),
            ('legal', '律师资格证怎么考', 2),
            ('legal', '离婚律师演员表', 1),
            ('ecommerce', '如何注册官方旗舰店', 2),
            ('ecommerce', '官方旗舰店如何注册', 2),
            ('ecommerce', '想开官方旗舰店需要多少钱', 2),
            ('ecommerce', '官方旗舰店装修教程', 2),
            ('ecommerce', '官方旗舰店换货电话', 5),
            ('ecommerce', '官方旗舰店购买', 10),
            ('ecommerce', '面试服装官方旗舰店购买', 10),
            ('ecommerce', '官方旗舰店客服面试经验', 2),
        ]
        for industry, term, expected in samples:
            with self.subTest(term=term):
                self.assertEqual(score_keyword(term, load_profile(industry))[0], expected)

    def test_custom_rule_compile_expansion_is_bounded(self):
        for pattern in ['a{1000000000}', '(a{1000}){1000}', r'a{1,1000000000}', 'a{,1001}']:
            with self.subTest(pattern=pattern):
                profile = {
                    'industry': '宠物洗护', 'conversion_goal': '预约洗护服务',
                    'rules': [{'pattern': pattern, 'score': 10, 'reason': '合成编译边界测试'}],
                    'fallback': {'score': 3, 'reason': '信息查询'},
                }
                with self.assertRaises(ValueError):
                    validate_profile(profile, '合成画像')


if __name__ == '__main__':
    unittest.main(verbosity=2)
