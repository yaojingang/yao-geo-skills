#!/usr/bin/env python3
"""Check synthetic Douyin paths, interruption handling, screenshots, and input-only gates."""
import base64
import copy
import hashlib
import json
from pathlib import Path
from tempfile import TemporaryDirectory
from openpyxl import load_workbook
from build_workbook import build, load_capture

PNG = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII=')
URL = 'https://www.douyin.com/jingxuan'


def query(term, depth, suggestions, status='ok'):
    item = {
        'query': term, 'depth': depth, 'status': status,
        'observed_input': term, 'captured_at': '2026-09-30T21:00:00+08:00',
        'page_url_before': URL, 'page_url': URL,
        'screenshots': [f'evidence/synthetic-{hashlib.sha256(term.encode()).hexdigest()[:16]}.png'],
        'suggestions': [{'term': word, 'rank': rank} for word, rank in suggestions],
    }
    if status == 'empty':
        item['empty_confirmation'] = '合成测试：已核对完整输入与无下拉状态'
    return item


def main():
    with TemporaryDirectory() as folder:
        base = Path(folder)
        (base / 'evidence').mkdir()
        (base / 'evidence/synthetic.png').write_bytes(PNG)
        data = {
            'platform': 'douyin', 'capture_mode': 'input_only', 'industry': 'study_abroad',
            'conversion_goal': '咨询或委托留学服务机构进行申请规划与办理',
            'roots': ['留学公司', '留学机构'],
            'queries': [
                query('留学公司', 0, [('留学费用', 1), ('留学咨询', 2), ('留学公司', 3)]),
                query('留学机构', 0, [('留学咨询', 1)]),
                query('留学咨询', 1, [('留学咨询机构', 1), ('留学公司', 2)]),
            ],
        }
        path = base / 'capture.json'
        out = base / 'keywords.xlsx'

        def save(value):
            for item in value['queries']:
                for shot in item.get('screenshots', []):
                    target = base / shot
                    if target.parent == base / 'evidence' and target.name.startswith('synthetic-'):
                        target.write_bytes(PNG)
            path.write_text(json.dumps(value, ensure_ascii=False), encoding='utf-8')

        def rejects(value, label):
            save(value)
            try:
                load_capture(path)
            except ValueError:
                return
            raise AssertionError('未拒绝：' + label)

        save(data)
        assert build(path, out) == (8, 4, 1)
        wb = load_workbook(out)
        assert wb.sheetnames == ['关键词表', '采集记录', '截图索引', '去重词库', '待采集']
        assert wb['待采集']['A2'].value == '留学费用'
        values = list(wb['关键词表'].values)[1:]
        assert sum(row[1] == 2 for row in values) == 4
        assert all(type(row[5]) is int and 1 <= row[5] <= 10 for row in values)
        assert len({row[5] for row in values if row[3] == '留学咨询机构'}) == 1
        assert wb['关键词表']['K2'].hyperlink.target == data['queries'][0]['screenshots'][0]
        assert wb['截图索引'].max_row == 4

        pending_only = copy.deepcopy(data)
        pending_only['queries'] = [{'query': '留学公司', 'depth': 0, 'status': 'pending'}]
        save(pending_only)
        assert build(path, out) == (0, 0, 2)

        spaced = copy.deepcopy(data)
        spaced['roots'][0] = ' 留学公司 '
        spaced['queries'][0]['query'] = ' 留学公司 '
        spaced['queries'][0]['observed_input'] = ' 留学公司 '
        spaced['queries'][0]['suggestions'][2]['term'] = ' 留学公司 '
        save(spaced)
        assert build(path, out)[0] == 8
        assert load_workbook(out)['关键词表']['A2'].value == ' 留学公司 '

        (base / 'evidence/second.png').write_bytes(PNG)
        multi = copy.deepcopy(data)
        multi['queries'][0]['screenshots'].append('evidence/second.png')
        for suggestion in multi['queries'][0]['suggestions']:
            suggestion['screenshot'] = 'evidence/second.png' if suggestion['rank'] == 2 else multi['queries'][0]['screenshots'][0]
        save(multi)
        build(path, out)
        assert load_workbook(out)['关键词表']['K3'].hyperlink.target == 'evidence/second.png'
        bad = copy.deepcopy(multi); bad['queries'][0]['suggestions'][0].pop('screenshot')
        rejects(bad, '多屏截图未逐词对应')
        bad = copy.deepcopy(multi); bad['queries'][0]['suggestions'][0]['screenshot'] = 'evidence/unlisted.png'
        rejects(bad, '建议词截图不属于查询')

        bad = copy.deepcopy(data); bad['queries'][0]['observed_input'] = '旧词'
        rejects(bad, '旧输入')
        bad = copy.deepcopy(data); bad['queries'][0]['page_url'] = 'https://www.douyin.com/search/test'
        rejects(bad, '意外跳转')
        bad = copy.deepcopy(data); bad['platform'] = 'xiaohongshu'
        rejects(bad, '错误平台')
        bad = copy.deepcopy(data); bad['capture_mode'] = 'submit_search'
        rejects(bad, '错误模式')
        bad = copy.deepcopy(data); bad['queries'][0]['screenshots'] = ['../outside.png']
        rejects(bad, '越界截图')
        (base / 'evidence/wrong.png').write_bytes(b'\xff\xd8\xff' + b'JPEG test bytes')
        bad = copy.deepcopy(data); bad['queries'][0]['screenshots'] = ['evidence/wrong.png']
        rejects(bad, 'JPEG伪装PNG')
        (base / 'evidence/truncated.png').write_bytes(b'\x89PNG\r\n\x1a\n')
        bad = copy.deepcopy(data); bad['queries'][0]['screenshots'] = ['evidence/truncated.png']
        rejects(bad, '只有签名的损坏截图')
        bad = copy.deepcopy(data); bad['queries'][0]['suggestions'][1]['rank'] = 1
        rejects(bad, '重复位置')
        bad = copy.deepcopy(data); bad['queries'].append(query('其他批次词', 1, [('其他结果', 1)]))
        rejects(bad, '没有一级来源的查询')
        bad = copy.deepcopy(data); bad['queries'][1]['screenshots'] = bad['queries'][0]['screenshots']
        rejects(bad, '不同检索共用截图文件')
        bad = copy.deepcopy(data); bad['queries'][0]['suggestions'][0]['term'] = 'x' * 32768
        rejects(bad, 'Excel会截断的超长词')

        empty = copy.deepcopy(data)
        empty['queries'].append(query('留学费用', 1, [], status='empty'))
        save(empty)
        assert build(path, out)[2] == 0
        bad = copy.deepcopy(empty); bad['queries'][-1].pop('empty_confirmation')
        rejects(bad, '加载未确认却记为空')

        blocked = copy.deepcopy(data)
        blocked['queries'].append({
            'query': '留学费用', 'depth': 1, 'status': 'blocked',
            'captured_at': '2026-09-30T21:01:00+08:00',
            'suggestions': [], 'screenshots': [], 'note': '合成测试：需用户登录',
        })
        save(blocked)
        assert build(path, out)[2] == 1
        assert load_workbook(out)['待采集']['B2'].value == 'blocked'

        duplicate = copy.deepcopy(data)
        duplicate['queries'][0]['suggestions'].append({'term': '留学咨询', 'rank': 4})
        save(duplicate)
        assert build(path, out)[0] == 11

        save(data)
        try:
            build(path, base / 'other/keywords.xlsx')
        except ValueError:
            pass
        else:
            raise AssertionError('应拒绝破坏截图相对链接的输出目录')
        try:
            build(path, path)
        except ValueError:
            pass
        else:
            raise AssertionError('应拒绝覆盖原始JSON')
        print('通过合成检查：两层路径、共享查询、自身词、空结果、受阻、评分及截图与输入模式校验')


if __name__ == '__main__':
    main()
