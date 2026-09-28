#!/usr/bin/env python3
"""Exercise workbook generation with duplicate terms and multiple screenshots."""

from __future__ import annotations

import argparse
import base64
import json
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

from build_workbook import build_workbook, load_items
from score_keywords import load_profile


PIXEL = base64.b64decode(
    'iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jR9sAAAAASUVORK5CYII='
)


def main() -> int:
    argparse.ArgumentParser(description=__doc__).parse_args()
    with TemporaryDirectory() as folder:
        run = Path(folder)
        evidence = run / 'evidence'
        evidence.mkdir()
        for name in ('001-autocomplete.png', '001-autocomplete-2.png', '001-bottom.png'):
            (evidence / name).write_bytes(PIXEL)

        capture = {
            'industry': 'legal',
            'items': [{
                'seed': '婚姻律师',
                'captured_at': '2026-09-28T12:00:00+08:00',
                'search_url': 'https://www.baidu.com/s?wd=%E5%A9%9A%E5%A7%BB%E5%BE%8B%E5%B8%88',
                'autocomplete': {
                    'status': 'ok',
                    'screenshots': ['evidence/001-autocomplete.png', 'evidence/001-autocomplete-2.png'],
                    'terms': ['婚姻律师事务所', '婚姻律师事务所', '离婚律师'],
                },
                'bottom': {
                    'status': 'ok',
                    'screenshots': ['evidence/001-bottom.png'],
                    'terms': ['婚姻律师事务所', '离婚律师推荐', '=1+1'],
                },
                'note': '=1+1',
            }],
        }
        source = run / 'capture.json'
        source.write_text(json.dumps(capture, ensure_ascii=False), encoding='utf-8')
        output = run / 'keywords.xlsx'
        rows, seeds = build_workbook(load_items(source), output, load_profile('legal'))
        assert (rows, seeds) == (4, 1)

        workbook = load_workbook(output)
        terms = workbook['关键词表']
        assert terms.max_row == 5
        assert terms['F1'].value == '商业价值评分（10分）'
        assert (terms['D2'].value, terms['E2'].value) == (1, 1)
        assert terms['C2'].value == '下拉框；页尾相关搜索'
        assert terms['I2'].value == 'evidence/001-autocomplete.png；evidence/001-autocomplete-2.png'
        assert terms['I2'].hyperlink.target == 'evidence/001-autocomplete.png'
        assert terms['J2'].hyperlink.target == 'evidence/001-bottom.png'
        assert all(type(terms.cell(row, 6).value) is int for row in range(2, 6))
        assert terms['B5'].value == '=1+1' and terms['B5'].data_type == 's'
        assert workbook['采集记录']['K2'].data_type == 's'

        index = workbook['截图索引']
        assert index.max_row == 4
        assert [index.cell(row, 4).hyperlink.target for row in range(2, 5)] == [
            'evidence/001-autocomplete.png',
            'evidence/001-autocomplete-2.png',
            'evidence/001-bottom.png',
        ]

        capture['items'][0]['search_url'] = ''
        source.write_text(json.dumps(capture, ensure_ascii=False), encoding='utf-8')
        try:
            load_items(source)
        except ValueError as exc:
            assert 'search_url 缺失' in str(exc)
        else:
            raise AssertionError('页尾相关搜索缺少结果页 URL 时必须拒绝生成')
    print('通过 Excel 去重、评分、文本安全和多截图索引检查')
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
