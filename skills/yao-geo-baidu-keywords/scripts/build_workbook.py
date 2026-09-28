#!/usr/bin/env python3
"""Turn browser-observed Baidu suggestions and screenshot evidence into an XLSX."""

from __future__ import annotations

import argparse
import json
import os
from datetime import datetime
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from score_keywords import ALIASES, load_profile, score_keyword


SECTIONS = (("autocomplete", "下拉框"), ("bottom", "页尾相关搜索"))
STATUSES = {"ok": "已采集", "empty": "无词条", "blocked": "页面受阻", "error": "采集失败"}
HEADER_FILL = PatternFill("solid", fgColor="17365D")
ALT_FILL = PatternFill("solid", fgColor="F2F6FC")


def text_value(value: object, label: str) -> str:
    if not isinstance(value, str):
        raise ValueError(f"{label} 必须是文本")
    result = " ".join(value.split())
    if not result:
        raise ValueError(f"{label} 不能为空")
    return result


def screenshot_paths(value: object, base: Path, label: str) -> list[Path]:
    if not isinstance(value, list) or not all(isinstance(path, str) for path in value):
        raise ValueError(f"{label}.screenshots 必须是路径数组")
    paths = []
    for raw in value:
        relative = Path(raw)
        resolved = (base / relative).resolve()
        if relative.is_absolute() or not resolved.is_relative_to(base):
            raise ValueError(f"{label} 截图必须位于采集文件夹内: {raw}")
        if resolved.suffix.lower() not in {".png", ".jpg", ".jpeg", ".webp"}:
            raise ValueError(f"{label} 截图格式无效: {raw}")
        if not resolved.is_file():
            raise ValueError(f"{label} 截图不存在: {raw}")
        paths.append(resolved)
    return paths


def load_items(input_path: Path) -> list[dict]:
    base = input_path.parent.resolve()
    data = json.loads(input_path.read_text(encoding="utf-8"))
    if not isinstance(data, dict) or not isinstance(data.get("items"), list) or not data["items"]:
        raise ValueError("JSON 必须包含非空 items 数组")

    parsed = []
    seen_seeds = set()
    for index, raw in enumerate(data["items"], 1):
        if not isinstance(raw, dict):
            raise ValueError(f"第 {index} 个目标词记录必须是对象")
        seed = text_value(raw.get("seed"), f"第 {index} 个 seed")
        if seed in seen_seeds:
            raise ValueError(f"目标词重复: {seed}")
        seen_seeds.add(seed)

        captured_at = text_value(raw.get("captured_at"), f"{seed}.captured_at")
        try:
            timestamp = datetime.fromisoformat(captured_at)
        except ValueError as exc:
            raise ValueError(f"{seed}.captured_at 不是 ISO 时间") from exc
        if timestamp.tzinfo is None:
            raise ValueError(f"{seed}.captured_at 必须包含时区")

        search_url = raw.get("search_url", "")
        if not isinstance(search_url, str):
            raise ValueError(f"{seed}.search_url 必须是文本")
        if search_url:
            url = urlparse(search_url)
            if url.scheme not in {"http", "https"} or not url.hostname or not (url.hostname == "baidu.com" or url.hostname.endswith(".baidu.com")):
                raise ValueError(f"{seed}.search_url 必须是百度页面 URL")

        sections = {}
        for key, display in SECTIONS:
            section = raw.get(key)
            if not isinstance(section, dict):
                raise ValueError(f"{seed}.{key} 缺失")
            status = section.get("status")
            if status not in STATUSES:
                raise ValueError(f"{seed}.{key}.status 无效")
            terms = section.get("terms")
            if not isinstance(terms, list):
                raise ValueError(f"{seed}.{key}.terms 必须是数组")
            terms = [text_value(term, f"{seed}.{key}.terms") for term in terms]
            paths = screenshot_paths(section.get("screenshots", []), base, f"{seed}.{key}")
            if status == "ok" and (not terms or not paths):
                raise ValueError(f"{seed}.{key} 已采集状态必须有词条和截图")
            if status != "ok" and terms:
                raise ValueError(f"{seed}.{key} 非已采集状态不能填写词条")
            sections[key] = {"status": status, "terms": terms, "paths": paths, "label": display}

        if sections["bottom"]["status"] == "ok" and not search_url:
            raise ValueError(f"{seed}.search_url 缺失：页尾相关搜索已采集时必须记录结果页 URL")
        note = raw.get("note", "")
        if note is None:
            note = ""
        if not isinstance(note, str):
            raise ValueError(f"{seed}.note 必须是文本")

        parsed.append({
            "seed": seed,
            "captured_at": captured_at,
            "search_url": search_url,
            "sections": sections,
            "note": note.strip(),
        })
    return parsed


def screenshot_cell(paths: list[Path], output_dir: Path) -> str:
    return "；".join(os.path.relpath(path, output_dir) for path in paths)


def put_text(cell, value: str) -> None:
    cell.value = value
    cell.data_type = "s"


def prepare_sheet(sheet, headers: list[str], widths: list[int]) -> None:
    sheet.append(headers)
    sheet.freeze_panes = "A2"
    for number, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(number)].width = width
    sheet.row_dimensions[1].height = 30
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = Font(name="PingFang SC", size=11, bold=True, color="FFFFFF")
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def finish_sheet(sheet) -> None:
    sheet.auto_filter.ref = f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}"
    for row in sheet.iter_rows(min_row=2):
        sheet.row_dimensions[row[0].row].height = 27
        for cell in row:
            cell.font = Font(name="PingFang SC", size=10, color="202A35")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if cell.row % 2 == 0:
                cell.fill = ALT_FILL


def build_workbook(items: list[dict], output_path: Path, profile: dict) -> tuple[int, int]:
    workbook = Workbook()
    terms_sheet = workbook.active
    terms_sheet.title = "关键词表"
    logs_sheet = workbook.create_sheet("采集记录")
    evidence_sheet = workbook.create_sheet("截图索引")

    prepare_sheet(
        terms_sheet,
        ["目标关键词", "相关搜索词", "来源", "下拉序号", "页尾序号", "商业价值评分（10分）", "采集时间", "搜索结果页", "下拉截图", "页尾截图"],
        [24, 42, 22, 12, 12, 22, 28, 48, 38, 38],
    )
    prepare_sheet(
        logs_sheet,
        ["目标关键词", "下拉状态", "下拉条数", "页尾状态", "页尾条数", "合并后条数", "采集时间", "搜索结果页", "下拉截图", "页尾截图", "备注"],
        [24, 14, 12, 14, 12, 15, 28, 48, 38, 38, 42],
    )
    prepare_sheet(evidence_sheet, ["目标关键词", "截图来源", "截图序号", "截图路径"], [24, 20, 12, 52])

    output_dir = output_path.parent.resolve()
    total_rows = 0
    for item in items:
        sections = item["sections"]
        merged = {}
        for source, _ in SECTIONS:
            for position, term in enumerate(sections[source]["terms"], 1):
                merged.setdefault(term, {}).setdefault(source, position)

        for source, label in SECTIONS:
            for position, path in enumerate(sections[source]["paths"], 1):
                relative = os.path.relpath(path, output_dir)
                evidence_sheet.append([item["seed"], label, position, relative])
                for column in (1, 2, 4):
                    put_text(evidence_sheet.cell(evidence_sheet.max_row, column), str(evidence_sheet.cell(evidence_sheet.max_row, column).value))

        autocomplete_path = screenshot_cell(sections["autocomplete"]["paths"], output_dir)
        bottom_path = screenshot_cell(sections["bottom"]["paths"], output_dir)
        for term, positions in merged.items():
            sources = "；".join(label for key, label in SECTIONS if key in positions)
            score, _ = score_keyword(term, profile)
            terms_sheet.append([
                item["seed"], term, sources,
                positions.get("autocomplete"), positions.get("bottom"),
                score,
                item["captured_at"], item["search_url"],
                autocomplete_path if "autocomplete" in positions else "",
                bottom_path if "bottom" in positions else "",
            ])
            for column in (1, 2, 3, 7, 8, 9, 10):
                put_text(terms_sheet.cell(terms_sheet.max_row, column), str(terms_sheet.cell(terms_sheet.max_row, column).value or ""))
            total_rows += 1

        logs_sheet.append([
            item["seed"], STATUSES[sections["autocomplete"]["status"]],
            len(sections["autocomplete"]["terms"]),
            STATUSES[sections["bottom"]["status"]],
            len(sections["bottom"]["terms"]), len(merged),
            item["captured_at"], item["search_url"],
            autocomplete_path, bottom_path, item["note"],
        ])
        for column in (1, 2, 4, 7, 8, 9, 10, 11):
            put_text(logs_sheet.cell(logs_sheet.max_row, column), str(logs_sheet.cell(logs_sheet.max_row, column).value or ""))

    for sheet, screenshot_columns in ((terms_sheet, (9, 10)), (logs_sheet, (9, 10)), (evidence_sheet, (4,))):
        finish_sheet(sheet)
        for row in sheet.iter_rows(min_row=2):
            for column in screenshot_columns:
                cell = row[column - 1]
                if cell.value:
                    cell.hyperlink = cell.value.split("；", 1)[0]
                    cell.font = Font(name="PingFang SC", size=10, color="1D5FBF", underline="single")

    workbook.properties.title = "百度相关搜索词关键词表"
    workbook.properties.subject = f"百度下拉框与页尾相关搜索采集；{profile['industry']}商业价值评分"
    workbook.properties.description = f"评分转化目标：{profile['conversion_goal']}；规则版本：{profile.get('version', 'custom')}"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return total_rows, len(items)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, type=Path, help="浏览器采集 JSON")
    parser.add_argument("--output", required=True, type=Path, help="输出 .xlsx 文件")
    parser.add_argument("--industry", help="行业：legal、medical、education、study_abroad、renovation、ecommerce 或中文别名")
    parser.add_argument("--profile", type=Path, help="其他行业或不同转化目标的自定义评分画像 JSON")
    args = parser.parse_args()
    if args.output.suffix.lower() != ".xlsx":
        parser.error("--output 必须以 .xlsx 结尾")
    try:
        input_path = args.input.resolve()
        output_path = args.output.resolve()
        if input_path.parent != output_path.parent:
            raise ValueError("采集 JSON 与 Excel 必须保存在同一运行文件夹")
        capture = json.loads(input_path.read_text(encoding="utf-8"))
        recorded_industry = capture.get("industry", "") if isinstance(capture, dict) else ""
        if recorded_industry and not isinstance(recorded_industry, str):
            raise ValueError("industry 必须是文本")
        if args.industry and recorded_industry and ALIASES.get(args.industry, args.industry) != ALIASES.get(recorded_industry, recorded_industry):
            raise ValueError("命令行行业与采集 JSON 中的 industry 不一致")
        industry = args.industry or recorded_industry
        profile = load_profile(industry, args.profile.resolve() if args.profile else None)
        recorded_goal = capture.get("conversion_goal", "") if isinstance(capture, dict) else ""
        if recorded_goal and recorded_goal != profile["conversion_goal"]:
            raise ValueError("采集 JSON 的 conversion_goal 与评分画像不一致；请使用匹配的 --profile")
        items = load_items(input_path)
        rows, seeds = build_workbook(items, output_path, profile)
    except (OSError, ValueError, json.JSONDecodeError) as exc:
        parser.exit(2, f"生成失败: {exc}\n")
    print(f"已生成 {args.output}：{seeds} 个目标词，{rows} 条合并后相关词；评分行业：{profile['industry']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
